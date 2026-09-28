"""Replace only S13-S19 illustrations and captions in the supplied DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(r"C:\Users\Google\Documents\HMSC-HIST")
SOURCE = Path(
    r"C:\Users\Google\Downloads"
    r"\HmscEE_MEE_Results_supplementary_renumbered_after_S2_removal.docx"
)
OUT_DIR = (
    ROOT
    / "outputs"
    / "HmscEcoEvo"
    / "manuscript_results_six_process_20260927"
    / "hmscHist_supplement_clean_20260928"
)
TARGET = OUT_DIR / "HmscEE_MEE_Results_supplementary_S13-S19_clean.docx"

CAPTIONS = {
    13: (
        "Derived environmental-response shape in a controlled software example. "
        "(A) Optimum and breadth summaries across the demonstration species and "
        "response axes. (B) Multivariate response directions summarized as "
        "vectors. The direct coefficient posterior and ordinary response-curve "
        "panels present in the original figure have been omitted. These are not "
        "Plant200 estimates or historical occupancy probabilities."
    ),
    14: (
        "Localization of phylogenetic structure in fitted responses in a controlled "
        "example. (A) Similarity across phylogenetic distances. (B) Scale-dependent "
        "signal and clade conservatism. (C) Tip-level local signal. (D) Internal-node "
        "signal. The original global HMSC rho summary and residual-coefficient heatmap "
        "were omitted; branch colours identify demonstration clades, not Plant200 taxa."
    ),
    15: (
        "Evolutionary-transition diagnostics from supplied demonstration histories. "
        "(A) Branch-level shift support and the distribution of total shifts. "
        "(B) Shift magnitude by environmental response axis. (C) Shifts and relative "
        "rates placed on the dated example tree. (D) Branch- and clade-level rate "
        "diagnostics. The display does not infer historical Plant200 shifts from "
        "modern occurrence records."
    ),
    16: (
        "Trait- and history-related explanation of response variation in a controlled "
        "example. (A) Descriptive indices for traits, historical predictors, their "
        "combined fit, and residual phylogenetic signal for each response axis; the "
        "indices are not additive fractions or causal mediation effects. (B) Trait fit "
        "versus residual signal. (C) Trait explanatory overlap versus phylogenetic "
        "signal. (D) Change in residual signal when each trait is omitted. The "
        "ordinary HMSC Gamma matrix and residual-coefficient display were omitted."
    ),
    17: (
        "Variation in trait-to-response associations in a controlled example. "
        "(A) Clade-specific Gamma patterns. (B) Regime-specific Gamma patterns. "
        "(C) Phylogenetic locations of Gamma-shift support. (D) Sign-reversal "
        "index and between-clade Gamma variance. These stratified diagnostics "
        "extend the single global HMSC Gamma display; they are not fitted "
        "Plant200 clade or regime effects."
    ),
    18: (
        "Population-level diagnostics using supplied demonstration data. "
        "(A) Within-species response divergence. (B) Descriptive partition of "
        "species-mean, local-adaptation, plasticity, and residual contributions. "
        "(C) Local-adaptation to plasticity ratio. (D) Genotype-specific reaction "
        "norms across an environmental gradient. The ordinary species-versus-population "
        "coefficient comparison was omitted. These displays require repeated "
        "population or genotype data absent from Plant200 and do not alone "
        "establish causal local adaptation."
    ),
    19: (
        "Robustness checks for derived hmscHist diagnostics in a controlled example. "
        "(A) Recovery of supplied true values for niche breadth, rate ratio, "
        "residual phylogenetic signal, and shift magnitude. (B) Sensitivity of derived "
        "metrics to tree uncertainty. (C) Sensitivity to priors, omitted environmental "
        "variables, and spatial structure. Standard HMSC block-validation, posterior "
        "predictive, and MCMC panels were omitted. This is not independent fossil or "
        "endpoint validation of Plant200."
    ),
}


def set_caption(paragraph, number: int, text: str) -> None:
    paragraph.clear()
    label = paragraph.add_run(f"Supplementary Figure S{number}. ")
    label.bold = True
    paragraph.add_run(text)
    paragraph.paragraph_format.keep_together = True
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9)


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)
    intro = next(
        p for p in doc.paragraphs
        if p.text.startswith("The following seven figures use controlled software inputs.")
    )
    intro.text = (
        "The following seven figures use controlled software inputs to demonstrate "
        "hmscHist diagnostics beyond standard HMSC displays. Direct HMSC coefficient "
        "and Gamma plots and routine cross-validation, posterior predictive, and MCMC "
        "panels have been removed; retained panels are relabelled in reading order. "
        "The analysed Plant200 HMSC did not fit measured traits, phylogenetic "
        "covariance or species-specific historical predictors. None of these seven "
        "figures is an empirical Plant200 effect."
    )

    replaced = []
    for number, caption in CAPTIONS.items():
        caption_paragraph = next(
            p for p in doc.paragraphs
            if p.text.startswith(f"Supplementary Figure S{number}.")
        )
        image_paragraph = caption_paragraph._element.getprevious()
        if image_paragraph is None or not image_paragraph.xpath(".//w:drawing"):
            raise ValueError(f"S{number}: expected original image immediately before caption")
        from docx.text.paragraph import Paragraph

        figure_paragraph = Paragraph(image_paragraph, caption_paragraph._parent)
        figure_paragraph.clear()
        figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure_paragraph.paragraph_format.keep_with_next = True
        figure_paragraph.paragraph_format.space_before = Pt(6)
        figure_paragraph.paragraph_format.space_after = Pt(3)
        figure = (
            OUT_DIR
            / "figures"
            / f"Supplementary_Figure_S{number}_hmscHist_clean.png"
        )
        if not figure.is_file():
            raise FileNotFoundError(figure)
        figure_paragraph.add_run().add_picture(str(figure), width=Inches(6.35))
        set_caption(caption_paragraph, number, caption)
        replaced.append(number)

    if replaced != list(range(13, 20)):
        raise AssertionError(f"Unexpected figures replaced: {replaced}")
    doc.core_properties.title = "HmscEE Supplementary Results with Clean hmscHist Figures"
    doc.save(TARGET)
    verify = Document(TARGET)
    figure_text = [p.text for p in verify.paragraphs if p.text.startswith("Supplementary Figure S")]
    if len(verify.inline_shapes) != 19 or len(figure_text) != 19:
        raise AssertionError(
            f"Expected 19 images and captions; got {len(verify.inline_shapes)} "
            f"images and {len(figure_text)} captions"
        )
    print(TARGET)
    print("Figures and captions:", len(verify.inline_shapes), len(figure_text))


if __name__ == "__main__":
    main()
