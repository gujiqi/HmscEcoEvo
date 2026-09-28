#!/usr/bin/env python3
"""Build the UTF-8-safe Chinese Case05 Word atlas from completed PNG maps.

The R runner deliberately writes numeric results, PNGs, GeoTIFFs, and a map
index only.  On this Windows R installation, officer cannot write Chinese text
to OOXML safely.  This companion uses python-docx so the computation and the
Chinese tutorial atlas remain reproducible and independent.
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


METRICS = [
    (
        "expected_sampled_surviving_lineage_richness",
        "期望采样现生存活谱系丰富度",
        "计算为同一格网中当时活跃谱系的边际占据概率之和。颜色越深表示情景 ensemble 中该格网预计具有更多活跃谱系。它不是全部历史植物物种丰富度，也不是化石记录的真实普查。",
    ),
    (
        "potential_relative_environmental_support_mass",
        "潜在相对环境支持总量",
        "计算为所有活跃谱系祖先环境支持度的总和，发生在扩散和局地存续之前。它是环境过滤层，不是概率；若高于最终占据丰富度，说明移动、定殖或存续提供额外限制。",
    ),
    (
        "mean_relative_environmental_support",
        "平均相对环境支持度",
        "由祖先环境 Beta 响应与同一时间的古环境相乘得到。它仅表示环境层的相对支持，不能直接写成历史出现概率或生态因果效应。",
    ),
    (
        "mean_arrival_probability_1myr",
        "平均一百万年到达概率",
        "由球面局地扩散、稀有长距离扩散与已占据来源共同生成。0 Ma 没有下一前向区间，该层为 NA，而不是伪造的过程值。该图不是观测到的迁徙事件。",
    ),
    (
        "mean_colonisation_probability_1myr",
        "平均一百万年定殖概率",
        "计算为到达后成功建立的概率。没有来源到达时定殖为零。定殖参数是过程情景参数，不是由单期出现背景资料估计出的历史定殖率。",
    ),
    (
        "mean_conditional_persistence_probability_1myr",
        "平均一百万年条件存续概率",
        "表示已经占据的局地种群在下一百万年继续存在的条件概率。它不是新物种出现机会，也不是完整谱系存续率。",
    ),
    (
        "mean_local_extinction_probability_1myr",
        "平均一百万年局地消失概率",
        "为条件局地存续的补数。它表示情景内的局地消失风险，不是全球谱系灭绝率，不能用于恢复未留下现生后代的灭绝历史。",
    ),
    (
        "mean_field_compositional_shannon_entropy",
        "Mean field Shannon 谱系多样性",
        "由边际占据概率组成的 Shannon 熵。它同时受谱系丰富度与均匀度影响，是模型占据概率的描述性摘要，不是丰度调查得到的正式 Shannon 多样性后验。",
    ),
    (
        "mean_field_compositional_gini_simpson",
        "Mean field Gini Simpson 谱系多样性",
        "由边际占据概率组成的 Gini Simpson 摘要。它强调谱系组成的均匀性，不是原位样方丰度的 Simpson 指数，也不是完整古植物群落多样性。",
    ),
]


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_text(doc: Document, text: str, style: str | None = None, *, size: float = 10.5) -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, size=size)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path,
                        help="Completed Case05 output directory.")
    parser.add_argument("--mode", choices=("full", "representative"), default="full")
    return parser.parse_args()


def read_map_index(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    if not rows:
        raise RuntimeError(f"Empty map index: {path}")
    missing = [row for row in rows if not Path(row["png"]).is_file()]
    if missing:
        raise RuntimeError(f"Map index contains missing PNGs; first: {missing[0]['png']}")
    return rows


def write_completion_word_path(output: Path, word_path: Path) -> None:
    completion = output / "06_summaries" / "case05_v5_run_completion.csv"
    if not completion.is_file():
        return
    with completion.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
        fieldnames = list(rows[0].keys()) if rows else []
    if "word_atlas" not in fieldnames:
        fieldnames.append("word_atlas")
    for row in rows:
        row["word_atlas"] = str(word_path.resolve()).replace("\\", "/")
    with completion.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    args = parse_args()
    output = args.output.resolve()
    index_path = output / "06_summaries" / "case05_v5_map_index.csv"
    rows = read_map_index(index_path)
    report_dir = output / "09_report"
    report_dir.mkdir(parents=True, exist_ok=True)
    word_path = report_dir / "Case05_Plant200_4deg_dynamic_earth_tutorial_zh.docx"

    by_metric: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        by_metric.setdefault(row["metric"], []).append(row)
    for metric_rows in by_metric.values():
        metric_rows.sort(key=lambda row: float(row["time_ma"]), reverse=True)

    representative_targets = (0, 20, 65, 100, 150, 200, 250, 300, 325)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Case05 Plant200 4度全球古地球动态结果图册")
    set_run_font(title_run, size=18, bold=True)
    add_text(doc,
        "本图册对应完整 4度古陆地格网的过程约束正向情景。现代 Plant200 记录以目标类群出现背景资料处理，未记录不等于确认缺失。因此图件用于检验模型结构、参数情景和时空模式，不是唯一真实古分布重建。",
        size=11)
    doc.add_heading("一 科学范围和逐步流程", level=1)
    add_text(doc,
        "1 现代 HMSC 在空间分块训练格网中拟合，使用目标类群记录强度作为现代观测协变量。2 从保留的 HMSC Beta 后验抽样沿定年树以 Brownian bridge 重建祖先环境响应。3 将每套祖先响应投影到对应时间的完整古环境格网。4 先以稳定 PALEOMAP carrier 完成板块携带，再用球面局地扩散和明确的稀有长距离扩散形成到达压力。5 到达与环境条件共同决定定殖；已占据格网依条件存续或局地消失。6 每个时间区间以 0.5 Myr 连续 CTMC 子步递推。")
    add_text(doc,
        "time_ma 的 Ma 值越大越古老。系统树根约为 325 Ma，因此 325 Ma 以前为树时间域之外，不应解释成零多样性。灰色格网为海洋或不可用生境。每个指标在全部时间片使用同一色标范围，故可在同一指标内部比较时间变化；不同指标的颜色深浅不可直接比较。")
    doc.add_heading("二 结果文件", level=1)
    add_text(doc, f"PNG 地图：{output / '07_png_maps'}")
    add_text(doc, f"GIS GeoTIFF：{output / '08_geotiff_maps'}")
    add_text(doc, f"地图索引：{index_path}")
    add_text(doc, f"图册模式：{args.mode}；地图总数：{len(rows)}。")
    doc.add_heading("三 逐指标全时间图册", level=1)

    image_count = 0
    for metric, title_text, explanation in METRICS:
        metric_rows = by_metric.get(metric, [])
        if not metric_rows:
            raise RuntimeError(f"Map index has no maps for metric {metric}")
        if image_count:
            doc.add_page_break()
        doc.add_heading(title_text, level=1)
        add_text(doc, explanation)
        lower, upper = metric_rows[0]["lower"], metric_rows[0]["upper"]
        add_text(doc, f"统一色标范围：[{lower}, {upper}]。")
        if args.mode == "representative":
            metric_rows = sorted(
                {min(metric_rows, key=lambda row, target=target: abs(float(row['time_ma']) - target))["time_ma"]: min(metric_rows, key=lambda row, target=target: abs(float(row['time_ma']) - target))
                 for target in representative_targets}.values(),
                key=lambda row: float(row["time_ma"]), reverse=True,
            )
        for row in metric_rows:
            add_text(doc,
                f"时间片：{row['time_ma']} Ma。Ma 越大越古老；灰色为海洋或不可用生境。",
                style="Heading 2", size=11)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(row["png"], width=Inches(6.55))
            image_count += 1

    doc.add_page_break()
    doc.add_heading("四 质量检查和解释边界", level=1)
    add_text(doc,
        "计算质量门要求每个有效时间片均有 PNG 和 GeoTIFF、每个指标色标固定、谱系丰富度不超过活跃谱系数量，并在每个时间区间先完成板块携带再计算主动扩散。最终验证表中现代零值为伪缺失、空间持留不独立于 HMSC 训练资料这两项会明确保留为 FALSE；这是科学限制而非软件错误。")
    add_text(doc,
        "因此，不得将本图册表述为唯一真实古分布、真实历史定殖率、真实局地灭绝率、完整植物群落丰富度或完全谱系灭绝史。可写入方法或结果的准确表述是：在所声明的现代出现背景观测设计、HMSC 后验、定年树、古环境、板块 carrier 和过程参数情景下，模型生成了采样现生存活谱系的网格尺度动态占据情景与不确定性摘要。")

    doc.save(word_path)
    write_completion_word_path(output, word_path)
    manifest = {
        "word_atlas": str(word_path.resolve()),
        "mode": args.mode,
        "images_embedded": image_count,
        "map_index": str(index_path.resolve()),
    }
    with (report_dir / "word_atlas_manifest.json").open("w", encoding="utf-8") as handle:
        json.dump(manifest, handle, ensure_ascii=False, indent=2)
    print(json.dumps(manifest, ensure_ascii=False))


if __name__ == "__main__":
    main()
