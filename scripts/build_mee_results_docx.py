"""Build the MEE Results or Supplementary Results from checked Markdown."""

from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(r"C:\Users\Google\Documents\HMSC-HIST")
OUT = ROOT / "outputs" / "HmscEcoEvo" / "manuscript_results_six_process_20260927"
SOURCE = Path(os.environ.get("HMSCEE_RESULTS_SOURCE", str(OUT / "RESULTS.md")))
TARGET = Path(os.environ.get(
    "HMSCEE_RESULTS_DOCX",
    str(OUT / "HmscEE_MEE_expanded_Results_with_figures.docx"),
))
DOC_TITLE = os.environ.get("HMSCEE_RESULTS_TITLE", "HmscEE Expanded Results")
DOC_SUBTITLE = os.environ.get(
    "HMSCEE_RESULTS_SUBTITLE", "Manuscript ready replacement section with source checked figures"
)
EXPECTED_FIGURES = os.environ.get("HMSCEE_EXPECTED_FIGURES")
MANUSCRIPT_MODE = os.environ.get("HMSCEE_MANUSCRIPT_MODE") == "1"


def add_inline(p, text: str) -> None:
    for chunk in re.split(r"(\*[^*]+\*)", text):
        if not chunk:
            continue
        if chunk.startswith("*") and chunk.endswith("*"):
            p.add_run(chunk[1:-1]).italic = True
        else:
            p.add_run(chunk.replace("`", ""))


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(11)
    p.paragraph_format.line_spacing = 1.08
    m = re.match(r"((?:Supplementary )?Figure (?:S\d+|\d+)\.)\s*(.*)", text)
    if m:
        label = p.add_run(m.group(1) + " ")
        label.bold = True
        label.font.size = Pt(9)
        label.font.name = "Arial"
        add_inline(p, m.group(2))
    else:
        add_inline(p, text)
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(61, 76, 86)


def add_figure(doc: Document, rel: str) -> None:
    image_path = OUT / rel
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    with Image.open(image_path) as img:
        px_width, px_height = img.size
    width = min(6.35, 6.8 * px_width / px_height)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.add_run().add_picture(str(image_path), width=Inches(width))


def main() -> None:
    content = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin = sec.bottom_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 47, 56)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.13
    for name, size, before, after in (("Heading 1", 17, 12, 7),
                                      ("Heading 2", 12, 12, 5)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(32, 68, 91)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    title_style = doc.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(21)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(35, 47, 56)
    title_border = title_style.element.get_or_add_pPr().find(qn("w:pBdr"))
    if title_border is not None:
        title_style.element.get_or_add_pPr().remove(title_border)
    subtitle_style = doc.styles["Subtitle"]
    subtitle_style.font.name = "Arial"
    subtitle_style.font.size = Pt(10)
    subtitle_style.font.color.rgb = RGBColor(82, 93, 101)

    if not MANUSCRIPT_MODE:
        title = doc.add_paragraph(style="Title")
        title.add_run(DOC_TITLE)
        subtitle = doc.add_paragraph()
        subtitle.add_run(DOC_SUBTITLE)
        subtitle.style = "Subtitle"
        note = doc.add_paragraph()
        note.add_run("Evidence scope  ").bold = True
        note.add_run(
            "Plant200 analyses are separated from bundled synthetic hmscHist "
            "demonstrations. Archived full 200-tip movement runs used package "
            "version 1.0.3; the six-process 1.0.4 package was checked and rerun "
            "on reduced terrestrial and full response/patch workflows."
        )

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", content) if p.strip()]
    image_count = 0
    for block in paragraphs:
        if block.startswith("# "):
            doc.add_paragraph(block[2:], style="Heading 1")
        elif block.startswith("## "):
            doc.add_paragraph(block[3:], style="Heading 2")
        elif block.startswith("!["):
            m = re.fullmatch(r"!\[[^]]*\]\(([^)]+)\)", block)
            if not m:
                raise ValueError(f"Malformed figure reference: {block}")
            add_figure(doc, m.group(1))
            image_count += 1
        elif block.startswith("Figure ") or block.startswith("Supplementary Figure "):
            add_caption(doc, block)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.keep_together = True
            add_inline(p, block.replace("\n", " "))

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if not MANUSCRIPT_MODE:
        footer.add_run("HmscEE  |  " + DOC_TITLE + "  |  Working manuscript insert")
    if EXPECTED_FIGURES is not None and image_count != int(EXPECTED_FIGURES):
        raise ValueError(f"Expected {EXPECTED_FIGURES} figures; found {image_count} in {SOURCE}")
    doc.save(TARGET)
    print(f"Saved {TARGET} with {image_count} figures")


if __name__ == "__main__":
    main()
